#!/usr/bin/env python3
"""Create CSV or XLSX templates for a baoyan information tracker."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path


SHEETS = {
    "院校库.csv": [
        "school_id",
        "学校",
        "学院/学部",
        "城市",
        "地区",
        "学校层次",
        "学科/专业",
        "信息年份",
        "是否往届信息",
        "学科排名",
        "专业类型",
        "学费",
        "学制",
        "是否有夏令营",
        "夏令营时间",
        "是否有预推免",
        "预推免时间",
        "是否通常九推",
        "九推时间",
        "学校官网",
        "研究生院/研招网",
        "学院官网",
        "招生专题/报名系统",
        "官方公众号",
        "备注",
        "信息状态",
        "最后核验日期",
    ],
    "院校画像.csv": [
        "profile_id",
        "school_id",
        "学校",
        "学院/项目",
        "信息年份",
        "是否往届信息",
        "学校官网",
        "研究生院/研招网",
        "学院官网",
        "专业方向匹配",
        "学术/学科优势",
        "导师/研究方向",
        "专硕培养特点",
        "学硕培养特点",
        "城市与行业资源",
        "就业与实习线索",
        "学费与生活成本",
        "住宿/奖助信息",
        "申请难度",
        "考核偏好",
        "夏令营价值",
        "预推免价值",
        "九推机会",
        "主要优势",
        "主要风险",
        "适合人群",
        "不适合情形",
        "推荐档位",
        "推荐理由摘要",
        "画像来源 URL",
        "信息状态",
        "最后核验日期",
    ],
    "项目节点.csv": [
        "program_id",
        "school_id",
        "年份",
        "是否往届信息",
        "环节",
        "项目名称",
        "报名开始",
        "报名截止",
        "报名入口",
        "材料要求",
        "考核形式",
        "地点",
        "入营/入围通知日期",
        "考核日期",
        "结果发布日期",
        "夏令营效力",
        "效力依据",
        "offer 类型",
        "专硕/学硕",
        "招生名额",
        "学校官网",
        "研究生院/研招网",
        "学院官网",
        "来源标题",
        "来源 URL",
        "发布日期",
        "抓取日期",
        "信息状态",
        "待办",
    ],
    "我的申请.csv": [
        "application_id",
        "program_id",
        "是否计划报名",
        "优先级",
        "匹配度",
        "简历投递状态",
        "报名状态",
        "初审状态",
        "是否入营",
        "是否参加夏令营",
        "夏令营结果",
        "是否拿到夏令营 offer",
        "预推免报名状态",
        "面试状态",
        "九推填报状态",
        "材料清单",
        "下一步",
        "截止提醒",
        "用户备注",
    ],
    "更新日志.csv": [
        "log_id",
        "日期",
        "学校",
        "环节",
        "类型",
        "是否往届信息",
        "变更前",
        "变更后",
        "来源 URL",
        "处理状态",
    ],
    "提醒队列.csv": [
        "reminder_id",
        "program_id",
        "提醒时间",
        "提醒类型",
        "提醒内容",
        "定时任务频率",
        "定时任务状态",
        "优先级",
        "是否已通知",
        "通知时间",
    ],
}

VALIDATIONS = {
    "院校库": {
        "学校层次": ["985", "211", "双一流", "双非", "中外合作", "未知"],
        "专业类型": ["学硕", "专硕", "均有", "未知"],
        "是否有夏令营": ["有", "无", "未确认", "历史有"],
        "是否有预推免": ["有", "无", "未确认", "历史有"],
        "是否通常九推": ["有", "无", "不稳定", "未知"],
        "是否往届信息": ["是", "否", "待核验"],
        "信息状态": ["已官方核验", "待核验", "历史参考", "冲突"],
    },
    "院校画像": {
        "是否往届信息": ["是", "否", "待核验"],
        "专业方向匹配": ["高", "中", "低", "待核验"],
        "申请难度": ["高", "中", "低", "待评估"],
        "夏令营价值": ["强", "中", "弱", "无", "待核验"],
        "预推免价值": ["强", "中", "弱", "无", "待核验"],
        "九推机会": ["高", "中", "低", "不稳定", "未知"],
        "推荐档位": ["冲刺", "稳妥", "保底", "观察", "不建议"],
        "信息状态": ["已官方核验", "待核验", "历史参考", "冲突"],
    },
    "项目节点": {
        "是否往届信息": ["是", "否"],
        "环节": ["夏令营", "预推免", "九推", "补录", "其他"],
        "考核形式": ["线上", "线下", "混合", "未公布"],
        "夏令营效力": ["强效力", "弱效力", "无效力", "未披露", "不适用"],
        "offer 类型": ["优秀营员", "候补", "拟录取", "待录取", "无"],
        "专硕/学硕": ["学硕", "专硕", "均有", "未知"],
        "信息状态": ["新增", "已核验", "有变更", "已截止", "历史参考", "冲突"],
    },
    "我的申请": {
        "是否计划报名": ["是", "否", "观望"],
        "优先级": ["S", "A", "B", "C"],
        "匹配度": ["高", "中", "低", "待评估"],
        "简历投递状态": ["未开始", "准备中", "已提交", "需补材料", "失败"],
        "报名状态": ["未报名", "填写中", "已报名", "已截止", "放弃"],
        "初审状态": ["未出", "通过", "未通过", "候补", "不适用"],
        "是否入营": ["是", "否", "未出", "不适用"],
        "是否参加夏令营": ["是", "否", "待定", "不适用"],
        "夏令营结果": ["优秀营员", "合格", "候补", "未通过", "未出", "不适用"],
        "是否拿到夏令营 offer": ["是", "否", "未出", "不适用"],
        "预推免报名状态": ["未报名", "已报名", "通过", "未通过", "不适用"],
        "面试状态": ["未安排", "待面试", "已面试", "通过", "未通过", "候补"],
        "九推填报状态": ["未填报", "已填报", "复试", "待录取", "已确认", "放弃"],
    },
    "更新日志": {
        "是否往届信息": ["是", "否", "不适用"],
        "类型": ["新增", "变更", "截止提醒", "名单发布", "无变化", "冲突"],
        "处理状态": ["未处理", "已更新表格", "需用户确认"],
    },
    "提醒队列": {
        "提醒类型": ["报名截止", "材料补交", "面试", "结果查询", "九推确认"],
        "定时任务频率": ["每日", "每周", "手动", "不设置"],
        "定时任务状态": ["待设置", "已设置", "暂停", "不适用"],
        "优先级": ["高", "中", "低"],
        "是否已通知": ["是", "否"],
    },
}


def write_csv(path: Path, headers: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(headers)


def write_xlsx(path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise SystemExit(
            "openpyxl is required for --format xlsx. Use --format csv or install openpyxl."
        ) from exc

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True, color="1F2937")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for filename, headers in SHEETS.items():
        sheet_name = filename.removesuffix(".csv")
        worksheet = workbook.create_sheet(title=sheet_name)
        worksheet.append(headers)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        for index, header in enumerate(headers, start=1):
            width = min(max(len(header) + 4, 12), 28)
            worksheet.column_dimensions[get_column_letter(index)].width = width

        for header, options in VALIDATIONS.get(sheet_name, {}).items():
            if header not in headers:
                continue
            column_letter = get_column_letter(headers.index(header) + 1)
            formula = '"' + ",".join(options) + '"'
            validation = DataValidation(type="list", formula1=formula, allow_blank=True)
            worksheet.add_data_validation(validation)
            validation.add(f"{column_letter}2:{column_letter}500")

    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create baoyan tracker CSV templates.")
    parser.add_argument("--out", default="baoyan-tracker", help="Output folder.")
    parser.add_argument(
        "--format",
        choices=["csv", "xlsx", "both"],
        default="csv",
        help="Template format to create.",
    )
    args = parser.parse_args()

    out_dir = Path(args.out).expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    if args.format in {"csv", "both"}:
        for filename, headers in SHEETS.items():
            write_csv(out_dir / filename, headers)

    if args.format in {"xlsx", "both"}:
        write_xlsx(out_dir / "保研追踪表模板.xlsx")

    print(f"Created baoyan tracker template files in {out_dir}")


if __name__ == "__main__":
    main()
